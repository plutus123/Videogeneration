"""Excel export utility — AlphaSense-grade report for AMCA intelligence.

Produces a styled .xlsx workbook with sheets:
  1. "Executive Summary"    — Report title + executive summary + section summaries
  2. "Curated Articles"     — All curated articles with citations, scores, key facts
  3. "Tavily Search Results" — Raw search data from Tavily
  4. "Citations"            — Full citation reference list (AlphaSense-style)
"""

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------
HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1A2744", end_color="1A2744", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

SECTION_FONT = Font(name="Calibri", bold=True, size=12, color="1A2744")
SECTION_FILL = PatternFill(start_color="E8EDF3", end_color="E8EDF3", fill_type="solid")

TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="1A2744")
SUBTITLE_FONT = Font(name="Calibri", bold=True, size=11, color="4A90B8")
BODY_FONT = Font(name="Calibri", size=10)

DATA_FONT = Font(name="Calibri", size=10)
DATA_ALIGN = Alignment(vertical="top", wrap_text=True)
URL_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

ROW_FILL_ALT = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

# Score-based fills
SCORE_HIGH = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Green
SCORE_MED = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")   # Yellow
SCORE_LOW = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")   # Red


def _style_header(ws, headers, col_widths):
    """Write header row with styling."""
    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def _write_data_cell(ws, row, col, value, is_url=False, font=None):
    """Write a styled data cell."""
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font or (URL_FONT if is_url else DATA_FONT)
    cell.alignment = DATA_ALIGN
    cell.border = THIN_BORDER
    if row % 2 == 0:
        cell.fill = ROW_FILL_ALT
    if is_url and value and isinstance(value, str) and value.startswith("http"):
        cell.hyperlink = value
    return cell


def _write_executive_summary_sheet(wb, curated_data):
    """Write the executive summary sheet (AlphaSense-style cover page)."""
    ws = wb.active
    ws.title = "Executive Summary"
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 100

    row = 1

    # Report title
    title = curated_data.get("report_title", "AMCA & Indian Defence Aerospace Intelligence Report")
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    row += 1

    # Date
    cell = ws.cell(row=row, column=1, value=f"Generated: {datetime.now().strftime('%d %B %Y')}")
    cell.font = Font(name="Calibri", size=10, color="808080")
    row += 2

    # Executive summary
    exec_summary = curated_data.get("executive_summary", "")
    if exec_summary:
        cell = ws.cell(row=row, column=1, value="EXECUTIVE SUMMARY")
        cell.font = SECTION_FONT
        cell.fill = SECTION_FILL
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1
        cell = ws.cell(row=row, column=1, value=exec_summary)
        cell.font = BODY_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row + 3, end_column=2)
        row += 5

    # Section summaries
    for section in curated_data.get("thematic_sections", []):
        section_title = section.get("section_title", "")
        section_summary = section.get("section_summary", "")
        num_articles = len(section.get("articles", []))

        cell = ws.cell(row=row, column=1, value=f"📑 {section_title} ({num_articles} articles)")
        cell.font = SUBTITLE_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1

        if section_summary:
            cell = ws.cell(row=row, column=1, value=section_summary)
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=2)
            row += 3

    # Total count
    total = sum(len(s.get("articles", [])) for s in curated_data.get("thematic_sections", []))
    row += 1
    cell = ws.cell(row=row, column=1, value=f"Total curated articles: {total}")
    cell.font = Font(name="Calibri", bold=True, size=11)

    return ws


def _write_curated_sheet(wb, curated_data):
    """Write curated articles to a sheet with thematic section grouping."""
    ws = wb.create_sheet("Curated Articles")

    headers = [
        "[#]", "Section", "Title", "URL", "Source Type", "Source",
        "Score", "Published", "Summary", "Key Facts",
        "Key Players", "AMCA Keywords"
    ]
    col_widths = [5, 30, 45, 40, 14, 18, 8, 14, 55, 50, 25, 25]
    _style_header(ws, headers, col_widths)

    row = 2

    if "thematic_sections" in curated_data:
        # AlphaSense-style with sections
        for section in curated_data.get("thematic_sections", []):
            for article in section.get("articles", []):
                key_facts = "\n".join(f"• {f}" for f in article.get("key_facts", []))
                key_players = ", ".join(article.get("key_players_mentioned", []))
                amca_keywords = ", ".join(article.get("amca_keywords_found", []))

                _write_data_cell(ws, row, 1, article.get("citation_number", ""))
                _write_data_cell(ws, row, 2, section.get("section_title", ""))
                _write_data_cell(ws, row, 3, article.get("title", "N/A"))
                _write_data_cell(ws, row, 4, article.get("url", ""), is_url=True)
                _write_data_cell(ws, row, 5, article.get("source_type", ""))
                _write_data_cell(ws, row, 6, article.get("source_name", ""))

                # Score with color
                score = article.get("importance_score", 0)
                score_cell = _write_data_cell(ws, row, 7, score)
                if score >= 8:
                    score_cell.fill = SCORE_HIGH
                elif score >= 5:
                    score_cell.fill = SCORE_MED
                else:
                    score_cell.fill = SCORE_LOW

                _write_data_cell(ws, row, 8, article.get("published_date", ""))
                _write_data_cell(ws, row, 9, article.get("summary", ""))
                _write_data_cell(ws, row, 10, key_facts)
                _write_data_cell(ws, row, 11, key_players)
                _write_data_cell(ws, row, 12, amca_keywords)
                row += 1
    else:
        # Legacy flat format
        for i, article in enumerate(curated_data.get("selected_articles", []), 1):
            key_facts = "\n".join(f"• {f}" for f in article.get("key_facts", []))
            key_players = ", ".join(article.get("key_players_mentioned", []))
            amca_keywords = ", ".join(article.get("amca_keywords_found", []))

            _write_data_cell(ws, row, 1, i)
            _write_data_cell(ws, row, 2, article.get("bucket", ""))
            _write_data_cell(ws, row, 3, article.get("title", "N/A"))
            _write_data_cell(ws, row, 4, article.get("url", ""), is_url=True)
            _write_data_cell(ws, row, 5, "")
            _write_data_cell(ws, row, 6, "")
            _write_data_cell(ws, row, 7, article.get("importance_score", 0))
            _write_data_cell(ws, row, 8, "")
            _write_data_cell(ws, row, 9, article.get("summary", ""))
            _write_data_cell(ws, row, 10, key_facts)
            _write_data_cell(ws, row, 11, key_players)
            _write_data_cell(ws, row, 12, amca_keywords)
            row += 1

    if row > 2:
        ws.auto_filter.ref = f"A1:L{row - 1}"
    return ws


def _write_tavily_sheet(wb, search_results):
    """Write raw Tavily search results."""
    ws = wb.create_sheet("Tavily Search Results")

    headers = [
        "#", "Title", "URL", "Domain", "Published Date",
        "Tavily Score", "Content Preview"
    ]
    col_widths = [5, 50, 45, 20, 22, 14, 70]
    _style_header(ws, headers, col_widths)

    for i, article in enumerate(search_results, 1):
        row = i + 1
        content = article.get("content", "")
        preview = content[:500].replace("\n", " ").strip()
        if len(content) > 500:
            preview += "..."

        _write_data_cell(ws, row, 1, i)
        _write_data_cell(ws, row, 2, article.get("title", "N/A"))
        _write_data_cell(ws, row, 3, article.get("url", ""), is_url=True)
        _write_data_cell(ws, row, 4, article.get("domain", ""))
        _write_data_cell(ws, row, 5, article.get("published_date", ""))
        _write_data_cell(ws, row, 6, round(article.get("score", 0), 4))
        _write_data_cell(ws, row, 7, preview)

    ws.auto_filter.ref = f"A1:G{len(search_results) + 1}"
    return ws


def _write_citations_sheet(wb, curated_data):
    """Write citations reference sheet (AlphaSense-style)."""
    ws = wb.create_sheet("Citations")

    headers = ["[#]", "Source Type", "Source Name", "Date", "Title", "URL"]
    col_widths = [6, 16, 22, 16, 60, 50]
    _style_header(ws, headers, col_widths)

    citations = curated_data.get("citations", [])
    if not citations:
        # Build from thematic_sections if citations not provided separately
        for section in curated_data.get("thematic_sections", []):
            for article in section.get("articles", []):
                citations.append({
                    "citation_number": article.get("citation_number", ""),
                    "source_type": article.get("source_type", "News"),
                    "source_name": article.get("source_name", ""),
                    "date": article.get("published_date", ""),
                    "title": article.get("title", ""),
                    "url": article.get("url", ""),
                })

    for i, cite in enumerate(citations, 1):
        row = i + 1
        _write_data_cell(ws, row, 1, cite.get("citation_number", i))
        _write_data_cell(ws, row, 2, cite.get("source_type", ""))
        _write_data_cell(ws, row, 3, cite.get("source_name", ""))
        _write_data_cell(ws, row, 4, cite.get("date", ""))
        _write_data_cell(ws, row, 5, cite.get("title", ""))
        _write_data_cell(ws, row, 6, cite.get("url", ""), is_url=True)

    if citations:
        ws.auto_filter.ref = f"A1:F{len(citations) + 1}"
    return ws


def export_to_excel(search_results, curated_data, output_path):
    """Export to AlphaSense-grade styled Excel workbook.

    Sheets:
      1. Executive Summary — Report title, exec summary, section summaries
      2. Curated Articles  — All curated articles with scores and key facts
      3. Tavily Search Results — Raw search data
      4. Citations — Full reference list

    Args:
        search_results: List of article dicts from Tavily search.
        curated_data:   Dict from curation agent (AlphaSense-style or legacy flat).
        output_path:    Path for the output .xlsx file.

    Returns:
        Path to the created Excel file.
    """
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    wb = Workbook()

    # Sheet 1: Executive Summary
    _write_executive_summary_sheet(wb, curated_data)

    # Sheet 2: Curated Articles
    _write_curated_sheet(wb, curated_data)

    # Sheet 3: Raw Tavily results
    _write_tavily_sheet(wb, search_results)

    # Sheet 4: Citations
    _write_citations_sheet(wb, curated_data)

    wb.save(output_path)
    wb.close()

    num_search = len(search_results)
    total_curated = sum(
        len(s.get("articles", []))
        for s in curated_data.get("thematic_sections", [])
    ) or len(curated_data.get("selected_articles", []))
    num_sections = len(curated_data.get("thematic_sections", []))

    print(f"  📊 Excel exported: {output_path}")
    print(f"     Sheet 1: Executive Summary ({num_sections} sections)")
    print(f"     Sheet 2: {total_curated} curated articles")
    print(f"     Sheet 3: {num_search} search results")
    print(f"     Sheet 4: Citations")

    return output_path
